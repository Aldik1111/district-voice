import csv
import io

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text
from sqlalchemy.exc import IntegrityError

from app.main import create_app


@pytest.fixture
def curated(client, valid_record, upload_csv, confirm):
    rows = [dict(valid_record, source_record_id=f"TEST-{index}", source="ДЕМО · Портал" if index % 2 else "ДЕМО · Чат", date=day, text=message)
            for index, (day, message) in enumerate([
                ("2026-09-01", "На улице Абая 15 нет холодной воды."),
                ("2026-09-02", "На улице Абая 15 отключили холодную воду."),
                ("2026-09-02", "На улице Абая 15 нет холодной воды."),
                ("2026-10-01", "На улице Абая 15 отсутствует холодная вода."),
            ])]
    confirm(client, upload_csv(client, rows))
    assert client.post("/api/analysis/run", json={"mode": "tfidf"}).json()["status"] == "completed"
    groups = client.get("/api/groups").json()
    case = next(group for group in groups if group["kind"] == "case" and group["message_count"] == 3)
    problem = next(group for group in groups if group["kind"] == "problem")
    return case, problem


def action(client, group, operation, **fields):
    return client.post(f"/api/groups/{group['id']}/actions", json={"action": operation, "version": group["version"], "reason": "Проверка демо-оператора", **fields})


def test_remove_recounts_parent_keeps_appeal_and_restart(client, database_url, curated):
    case, problem = curated
    evidence = client.get(f"/api/groups/{case['id']}/appeals").json()
    removed = next(item for item in evidence if item["source_record_id"] == "TEST-0")
    response = action(client, case, "remove", appeal_ids=[removed["id"]])
    assert response.status_code == 200, response.text
    group = response.json()["group"]
    assert group["message_count"] == 2 and group["start_date"] == "2026-09-02"
    assert group["source_count"] == 2 and group["explanation_stale"]
    assert group["minimum_similarity"] is None
    assert client.get(f"/api/groups/{problem['id']}").json()["message_count"] == 3
    assert client.get("/api/registry", params={"q": removed["id"]}).json()["total"] == 1
    event = group["review_history"][-1]
    assert event["actor"] == "demo_operator" and event["action"] == "remove"
    assert removed["id"] in event["old_value"]["appeal_ids"] and removed["id"] not in event["new_value"]["appeal_ids"]
    assert action(client, case, "remove", appeal_ids=[removed["id"]]).status_code == 409
    with TestClient(create_app(database_url)) as restarted:
        assert restarted.get(f"/api/groups/{case['id']}").json() == group
        restored = action(restarted, group, "add", appeal_ids=[removed["id"]])
        assert restored.status_code == 200, restored.text
        assert restored.json()["group"]["message_count"] == 3
        assert restarted.get(f"/api/groups/{problem['id']}").json()["message_count"] == 4


def test_split_merge_preserve_ids_and_case_graph(client, curated):
    case, problem = curated
    chosen = case["appeal_ids"][:1]
    response = action(client, case, "split", appeal_ids=chosen)
    assert response.status_code == 200, response.text
    left, right = response.json()["group"], response.json()["created_group"]
    assert set(left["appeal_ids"]) | set(right["appeal_ids"]) == set(case["appeal_ids"])
    assert not set(left["appeal_ids"]) & set(right["appeal_ids"])
    parent = client.get(f"/api/groups/{problem['id']}").json()
    assert len(parent["case_ids"]) == 3 and parent["message_count"] == 4
    assert client.get("/api/analysis/runs").json()[0]["case_count"] == 3
    assert action(client, left, "merge", target_id=right["id"], target_version=1).status_code == 409
    merged = action(client, left, "merge", target_id=right["id"], target_version=right["version"])
    assert merged.status_code == 200, merged.text
    assert merged.json()["group"]["appeal_ids"] == case["appeal_ids"]
    assert client.get(f"/api/groups/{right['id']}").json()["status"] == "archived"
    assert right["id"] not in {item["id"] for item in client.get("/api/groups").json()}
    parent = client.get(f"/api/groups/{problem['id']}").json()
    assert len(parent["case_ids"]) == 2 and parent["message_count"] == 4
    assert client.get("/api/dashboard").json()["groups_count"] == 3
    assert client.get("/api/analysis/runs").json()[0]["case_count"] == 2


def test_problem_split_merge_and_direct_remove_add(client, curated):
    _, problem = curated
    reply = action(client, problem, "split", case_ids=problem["case_ids"][:1])
    assert reply.status_code == 200, reply.text
    left, right = reply.json()["group"], reply.json()["created_group"]
    merged = action(client, left, "merge", target_id=right["id"], target_version=right["version"]).json()["group"]
    assert merged["appeal_ids"] == problem["appeal_ids"]
    case_id = merged["case_ids"][0]
    appeal_id = client.get(f"/api/groups/{case_id}").json()["appeal_ids"][0]
    removed = action(client, merged, "remove", appeal_ids=[appeal_id]).json()["group"]
    assert removed["message_count"] == 3
    added = action(client, removed, "add", appeal_ids=[appeal_id], case_id=case_id)
    assert added.status_code == 200, added.text
    assert added.json()["group"]["message_count"] == 4


def test_edit_and_invalid_operations_rollback(client, curated):
    case, problem = curated
    confirmed = client.patch(f"/api/groups/{case['id']}/review", json={"version": case["version"], "status": "confirmed", "reason": "Состав проверен"}).json()
    updated = action(client, confirmed, "edit", title="Перебои подачи воды", category="Другое").json()["group"]
    assert updated["status"] == "proposed" and updated["title"] == "Перебои подачи воды"
    assert updated["category"] == "Другое"
    assert action(client, updated, "add", appeal_ids=["unknown"]).status_code == 422
    assert action(client, updated, "add", appeal_ids=case["appeal_ids"][:1]).status_code == 409
    assert action(client, updated, "split", appeal_ids=case["appeal_ids"]).status_code == 422
    assert action(client, updated, "merge", target_id=problem["id"], target_version=client.get(f"/api/groups/{problem['id']}").json()["version"]).status_code == 422
    assert action(client, updated, "edit", title="Телефон +7 777 123 45 67").status_code == 422
    assert client.get(f"/api/groups/{case['id']}").json() == updated


def draft(client, curated):
    _, problem = curated
    response = client.patch(f"/api/groups/{problem['id']}/review", json={"version": problem["version"], "status": "confirmed", "reason": "Проверка повторяемости"})
    assert response.status_code == 200
    scope = {"start": "2026-09-01", "end": "2026-10-31", "territory": "абай 15"}
    preview = client.post("/api/summaries/preview", json=scope).json()
    assert preview["total_messages"] == 4 and preview["confirmed_problems"] == 1
    assert preview["unique_messages"] == 3
    return client.post("/api/summaries", json=scope).json()


def test_approve_snapshot_export_immutable_across_later_changes(client, database_url, curated):
    item = draft(client, curated)
    path = f"/api/summaries/{item['id']}"
    assert client.get(path + "/export/csv").status_code == 409
    edited = client.patch(path, json={"revision": 1, "title": "=SUM(A1)", "body": "Синтетическая демонстрация. Основания проверены."}).json()
    approved = client.post(path + "/approve", json={"revision": edited["revision"]})
    assert approved.status_code == 200, approved.text
    approved = approved.json()
    assert approved["approved_at"] and approved["actor"] == "demo_operator"
    export = client.get(path + "/export/csv")
    rows = list(csv.DictReader(io.StringIO(export.content.decode("utf-8-sig"))))
    assert len(rows) == 6
    assert rows[0]["Название / текст"].startswith("'=SUM")
    assert all(row["Версия сводки"] == "1" and row["Время утверждения UTC"] == approved["approved_at"] for row in rows)
    assert set(rows[1]["Подтверждающие ID"].split("; ")) == set(curated[1]["appeal_ids"])
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(client.get(path + "/export/xlsx").content))
    assert workbook.active.max_row == 7
    assert not any(cell.data_type == "f" for row in workbook.active for cell in row)
    assert client.patch(path, json={"revision": approved["revision"], "title": "Изменено", "body": "Текст изменён"}).status_code == 409
    assert client.post(path + "/approve", json={"revision": approved["revision"]}).status_code == 409
    # A real group change cannot rewrite an approved report or its exported bytes.
    case = curated[0]
    assert action(client, case, "remove", appeal_ids=case["appeal_ids"][:1]).status_code == 200
    assert client.get(path + "/export/csv").content == export.content
    with TestClient(create_app(database_url)) as restarted:
        assert restarted.get(path).json() == approved
        assert restarted.get(path + "/export/csv").content == export.content
    for operation in ("UPDATE summaries SET body='changed'", "DELETE FROM summaries"):
        with pytest.raises(IntegrityError), client.app.state.engine.begin() as connection:
            connection.execute(text(operation + " WHERE id=:id"), {"id": item["id"]})
    successor = client.post("/api/summaries", json={**item["scope"], "previous_id": item["id"]}).json()
    assert successor["series_id"] == approved["series_id"] and successor["version"] == 2 and successor["status"] == "draft"


def test_stale_draft_and_scope_validation(client, curated):
    item = draft(client, curated)
    case = curated[0]
    action(client, case, "remove", appeal_ids=case["appeal_ids"][:1])
    assert client.post(f"/api/summaries/{item['id']}/approve", json={"revision": 1}).status_code == 409
    assert client.post("/api/summaries/preview", json={"start": "2026-10-01", "end": "2026-09-01"}).status_code == 422
    scoped = client.post("/api/summaries/preview", json={"start": "2026-10-01", "end": "2026-10-31"}).json()
    assert scoped["total_messages"] == 1 and scoped["unique_messages"] == 1
    empty = client.post("/api/summaries", json={"start": "2020-01-01", "end": "2020-02-01"}).json()
    assert client.post(f"/api/summaries/{empty['id']}/approve", json={"revision": 1}).status_code == 422


def test_rejected_child_invalidates_parent_and_cannot_be_approved_as_evidence(client, curated):
    case, problem = curated
    report = draft(client, curated)
    before = client.get(f"/api/groups/{problem['id']}").json()
    assert before["status"] == "confirmed"
    rejected = client.patch(f"/api/groups/{case['id']}/review", json={"version": case["version"], "status": "rejected", "reason": "Случай не подтверждён"})
    assert rejected.status_code == 200
    parent = client.get(f"/api/groups/{problem['id']}").json()
    assert parent["status"] == "proposed"
    assert parent["version"] == before["version"] + 1
    assert parent["appeal_ids"] == before["appeal_ids"]
    assert parent["review_history"][-1]["action"] == "related_review"
    assert client.post(f"/api/summaries/{report['id']}/approve", json={"revision": report["revision"]}).status_code == 409
    assert client.patch(f"/api/groups/{parent['id']}/review", json={"version": parent["version"], "status": "confirmed", "reason": "Повторная проверка"}).status_code == 409


@pytest.mark.parametrize("value", ["Текст с\x00символом", "Текст с\x0bсимволом", "Текст с\uffffсимволом"])
def test_invalid_text_edits_cannot_break_approved_exports(client, curated, value):
    item = draft(client, curated)
    path = f"/api/summaries/{item['id']}"
    assert client.patch(path, json={"revision": 1, "title": item["title"], "body": value}).status_code == 422
    assert action(client, curated[0], "edit", title=value).status_code == 422
    evidence = client.get(f"/api/groups/{curated[0]['id']}/appeals").json()[0]
    assert client.patch(f"/api/appeals/{evidence['id']}/privacy", json={"version": evidence["version"], "text": value, "address": evidence["address"]}).status_code == 422
    assert client.get(path).json() == item
    assert client.post(path + "/approve", json={"revision": 1}).status_code == 200
    assert client.get(path + "/export/xlsx").status_code == 200
