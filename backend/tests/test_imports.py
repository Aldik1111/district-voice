"""Behavior checks for the real preview -> confirm -> persisted registry flow."""

import io
import json
import sqlite3
import tempfile
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import create_app


def assert_counts(result, *, accepted=0, rejected=0, duplicates=0, conflicts=0):
    assert {key: result[key] for key in ("accepted", "rejected", "duplicates", "conflicts")} == {
        "accepted": accepted,
        "rejected": rejected,
        "duplicates": duplicates,
        "conflicts": conflicts,
    }


@pytest.mark.parametrize("bom", [False, True], ids=["utf8", "utf8-bom"])
def test_csv_preserves_cyrillic_and_only_persists_after_confirm(client, valid_record, upload_csv, confirm, bom):
    record = {**valid_record, "text": "Абай көшесінде су жоқ. Уличное освещение не работает."}
    preview = upload_csv(client, [record], bom=bom)
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows"][0]["text"] == record["text"]
    assert client.get("/api/appeals").json()["total"] == 0

    result = confirm(client, preview)
    assert_counts(result, accepted=1)
    response = client.get("/api/appeals").json()
    assert response["total"] == 1
    assert response["items"][0]["text"] == record["text"]
    assert response["items"][0]["source_record_id"] == record["source_record_id"]


def test_explicit_column_mapping_for_custom_headers(client, valid_record, upload_csv, confirm):
    mapping = {
        "source_record_id": "Номер сообщения",
        "source": "Площадка",
        "date": "Когда создано",
        "text": "Суть обращения",
        "address": "Место события",
    }
    records = [{header: valid_record[field] for field, header in mapping.items()}]
    result = confirm(client, upload_csv(client, records, mapping=mapping))
    assert_counts(result, accepted=1)
    saved = client.get("/api/appeals").json()["items"][0]
    for field in mapping:
        assert saved[field] == valid_record[field]


def test_xlsx_real_workbook_and_russian_headers(client, valid_record, confirm):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ID", "Источник", "Дата", "Текст", "Адрес"])
    sheet.append(list(valid_record.values()))
    output = io.BytesIO()
    workbook.save(output)
    preview = client.post(
        "/api/imports/preview",
        files={"file": ("обращения.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert_counts(confirm(client, preview), accepted=1)
    saved = client.get("/api/appeals").json()["items"][0]
    assert saved["text"] == valid_record["text"]
    assert saved["date"] == valid_record["date"]


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_blank_rows_do_not_shift_reported_error_row(client, valid_record, file_format):
    if file_format == "csv":
        content = (
            "source_record_id,source,date,text,address\n"
            '1,Портал,2026-10-01,"Первая строка\nвторая строка",Абая\n'
            "\n"
            "2,Портал,invalid,Нет воды,Абая\n"
        ).encode("utf-8")
        expected_row = 5
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(valid_record))
        sheet.append(list(valid_record.values()))
        sheet.append([])
        sheet.append(["KZ-002", "Портал", "invalid", "Нет воды", "Абая"])
        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
        expected_row = 4
    response = client.post("/api/imports/preview", files={"file": (f"rows.{file_format}", content)})
    assert response.status_code == 200, response.text
    invalid_rows = [row for row in response.json()["rows"] if row["errors"]]
    assert len(invalid_rows) == 1
    assert invalid_rows[0]["row_number"] == expected_row


def test_xlsx_formula_is_rejected_without_evaluation(client, valid_record):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(valid_record))
    sheet.append(["KZ-001", "Портал", "2026-10-01", '=HYPERLINK("https://example.invalid","Нет воды")', "Абая"])
    output = io.BytesIO()
    workbook.save(output)
    response = client.post("/api/imports/preview", files={"file": ("formula.xlsx", output.getvalue())})
    assert 400 <= response.status_code < 500, response.text
    assert response.json()["detail"]
    assert client.get("/api/appeals").json()["total"] == 0


@pytest.mark.parametrize(
    ("field", "value"),
    [("date", "2026-02-30"), ("date", "bad-date"), ("date", ""), ("text", ""), ("source", ""), ("source_record_id", "")],
)
def test_invalid_row_rejected_with_row_error(client, valid_record, upload_csv, confirm, field, value):
    preview = upload_csv(client, [{**valid_record, field: value}])
    assert preview.status_code == 200, preview.text
    assert preview.json()["invalid_rows"] == 1
    assert preview.json()["rows"][0]["errors"]
    assert_counts(confirm(client, preview), rejected=1)
    assert client.get("/api/appeals").json()["total"] == 0


def test_invalid_rows_do_not_prevent_valid_rows_in_same_batch(client, valid_record, upload_csv, confirm):
    invalid = {**valid_record, "source_record_id": "KZ-002", "date": "нет даты"}
    result = confirm(client, upload_csv(client, [valid_record, invalid]))
    assert_counts(result, accepted=1, rejected=1)
    assert client.get("/api/appeals").json()["total"] == 1


def test_repeat_import_and_conflict_do_not_overwrite_saved_record(client, valid_record, upload_csv, confirm):
    assert_counts(confirm(client, upload_csv(client, [valid_record])), accepted=1)
    assert_counts(confirm(client, upload_csv(client, [valid_record])), duplicates=1)
    changed = {**valid_record, "text": "Совершенно другое сообщение о ремонте дороги."}
    preview = upload_csv(client, [changed])
    assert preview.status_code == 200, preview.text
    assert preview.json()["conflict_rows"] == 1
    assert_counts(confirm(client, preview), conflicts=1)
    saved = client.get("/api/appeals").json()
    assert saved["total"] == 1
    assert saved["items"][0]["text"] == valid_record["text"]
    assert client.get("/api/stats").json()["total_appeals"] == 1


def test_same_text_with_new_id_or_source_is_a_separate_appeal(client, valid_record, upload_csv, confirm):
    records = [
        valid_record,
        {**valid_record, "source_record_id": "KZ-002"},
        {**valid_record, "source": "Районная приёмная"},
    ]
    assert_counts(confirm(client, upload_csv(client, records)), accepted=3)
    assert client.get("/api/appeals").json()["total"] == 3


def test_duplicates_and_conflicts_inside_one_file(client, valid_record, upload_csv, confirm):
    conflict = {**valid_record, "text": "Другая жалоба с тем же номером."}
    result = confirm(client, upload_csv(client, [valid_record, valid_record, conflict]))
    assert_counts(result, accepted=1, duplicates=1, conflicts=1)
    assert client.get("/api/appeals").json()["total"] == 1


def test_confirm_rechecks_identity_when_another_preview_was_confirmed(client, valid_record, upload_csv, confirm):
    first = upload_csv(client, [valid_record])
    second = upload_csv(client, [valid_record])
    changed = upload_csv(client, [{**valid_record, "text": "Изменённый текст до подтверждения."}])
    assert_counts(confirm(client, first), accepted=1)
    assert_counts(confirm(client, second), duplicates=1)
    assert_counts(confirm(client, changed), conflicts=1)
    assert client.get("/api/appeals").json()["total"] == 1


def test_confirm_is_idempotent_and_data_survives_new_app(client, valid_record, upload_csv, confirm, database_url):
    preview = upload_csv(client, [valid_record])
    first = confirm(client, preview)
    second = confirm(client, preview)
    assert first == second
    assert_counts(first, accepted=1)
    with TestClient(create_app(database_url=database_url)) as restarted:
        assert restarted.get("/api/appeals").json()["total"] == 1
        stats = restarted.get("/api/stats").json()
        assert stats["total_appeals"] == 1
        assert stats["total_batches"] == 1
        assert stats["by_source"] == [{"name": valid_record["source"], "count": 1}]
        assert stats["by_date"] == [{"date": valid_record["date"], "count": 1}]
        assert sum(row["count"] for row in stats["by_territory"]) == 1
        batches = restarted.get("/api/imports")
        assert batches.status_code == 200
        assert str(first["batch_id"]) in json.dumps(batches.json(), ensure_ascii=False)


def test_pagination_reports_global_count(client, valid_record, upload_csv, confirm):
    records = [{**valid_record, "source_record_id": f"KZ-{index}"} for index in range(5)]
    confirm(client, upload_csv(client, records))
    page1 = client.get("/api/appeals", params={"page": 1, "page_size": 2}).json()
    page2 = client.get("/api/appeals", params={"page": 2, "page_size": 2}).json()
    assert page1["total"] == page2["total"] == 5
    assert len(page1["items"]) == len(page2["items"]) == 2
    assert {row["id"] for row in page1["items"]}.isdisjoint(row["id"] for row in page2["items"])


@pytest.mark.parametrize("query", ["ОСВЕЩЕНИЕ", "КӨШЕСІ"])
def test_search_is_case_insensitive_for_russian_and_kazakh(client, valid_record, upload_csv, confirm, query):
    target = {**valid_record, "text": "Не работает освещение. Абай көшесі қараңғы."}
    other = {**valid_record, "source_record_id": "KZ-002", "text": "На остановке нет расписания."}
    confirm(client, upload_csv(client, [target, other]))
    response = client.get("/api/appeals", params={"q": query})
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 1
    assert response.json()["items"][0]["source_record_id"] == target["source_record_id"]


@pytest.mark.parametrize(
    ("filename", "content", "content_type"),
    [
        ("appeals.txt", b"not a supported file", "text/plain"),
        ("appeals.xlsx", b"not an xlsx archive", "application/octet-stream"),
        ("appeals.csv", b"", "text/csv"),
    ],
)
def test_malformed_or_unsupported_file_returns_clear_error(client, filename, content, content_type):
    response = client.post("/api/imports/preview", files={"file": (filename, content, content_type)})
    assert 400 <= response.status_code < 500, response.text
    assert response.json()["detail"]


def test_invalid_mapping_json_is_rejected(client):
    response = client.post(
        "/api/imports/preview",
        files={"file": ("appeals.csv", b"one,two\nx,y\n", "text/csv")},
        data={"mapping": "{not json}"},
    )
    assert 400 <= response.status_code < 500, response.text
    assert response.json()["detail"]


def test_unmapped_required_columns_are_shown_as_row_errors(client, valid_record, upload_csv, confirm):
    row = {"Неизвестный столбец": valid_record["text"]}
    preview = upload_csv(client, [row])
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["columns"] == ["Неизвестный столбец"]
    assert body["invalid_rows"] == 1
    assert body["rows"][0]["errors"]
    assert_counts(confirm(client, preview), rejected=1)


def test_unknown_preview_cannot_create_import(client):
    response = client.post("/api/imports/confirm", json={"preview_id": "does-not-exist"})
    assert response.status_code == 404, response.text
    assert client.get("/api/stats").json()["total_appeals"] == 0


def test_expired_preview_requires_reupload(client, valid_record, upload_csv, database_url):
    preview = upload_csv(client, [valid_record])
    assert preview.status_code == 200, preview.text
    preview_id = preview.json()["preview_id"]
    expired = (datetime.now(timezone.utc) - timedelta(hours=25)).replace(tzinfo=None)
    with sqlite3.connect(database_url.removeprefix("sqlite:///")) as connection:
        connection.execute("UPDATE import_previews SET created_at = ? WHERE id = ?", (expired.isoformat(" "), preview_id))
    response = client.post("/api/imports/confirm", json={"preview_id": preview_id})
    assert response.status_code == 410, response.text
    stats = client.get("/api/stats").json()
    assert stats["total_appeals"] == stats["total_batches"] == 0


def test_file_size_limit_is_enforced(client):
    response = client.post(
        "/api/imports/preview",
        files={"file": ("large.csv", b"x" * (5 * 1024 * 1024 + 1), "text/csv")},
    )
    assert 400 <= response.status_code < 500, response.text
    assert response.json()["detail"]


def test_chunked_multipart_without_content_length_is_bounded(client):
    def body():
        yield b'--test-boundary\r\nContent-Disposition: form-data; name="file"; filename="large.csv"\r\nContent-Type: text/csv\r\n\r\n'
        for _ in range(6):
            yield b"x" * (1024 * 1024)
        yield b"\r\n--test-boundary--\r\n"

    request = client.build_request(
        "POST", "/api/imports/preview", content=body(),
        headers={"Content-Type": "multipart/form-data; boundary=test-boundary"},
    )
    assert "content-length" not in request.headers
    response = client.send(request)
    assert response.status_code == 413, response.text
    assert response.json()["detail"]
    assert client.get("/api/stats").json()["total_appeals"] == 0


def test_accepted_upload_over_one_megabyte_never_spools_original_to_disk(client, valid_record, upload_csv, monkeypatch):
    def refuse_disk_spooling(_file):
        raise AssertionError("Original upload must stay in memory")

    monkeypatch.setattr(tempfile.SpooledTemporaryFile, "rollover", refuse_disk_spooling)
    records = [{**valid_record, "source_record_id": f"KZ-{index}", "text": "x" * 12_000} for index in range(100)]
    response = upload_csv(client, records)
    assert response.status_code == 200, response.text
    assert response.json()["valid_rows"] == 100


def test_row_limit_is_enforced(client, valid_record, upload_csv):
    records = [{**valid_record, "source_record_id": f"KZ-{index}"} for index in range(5001)]
    response = upload_csv(client, records)
    assert 400 <= response.status_code < 500, response.text
    assert response.json()["detail"]
