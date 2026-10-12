import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime

import pandas as pd
from fastapi import HTTPException
from openpyxl import load_workbook

from .privacy import detect_language, mask
from .input_validation import UNSUPPORTED_CHARACTERS

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 60
MAX_TEXT = 20_000
PREVIEW_TTL_HOURS = 24
FIELDS = ("source", "source_record_id", "date", "text", "address", "official_status")
REQUIRED = ("source", "source_record_id", "date", "text")
LABELS = {"source": "источник", "source_record_id": "ID в источнике", "date": "дата", "text": "текст", "address": "адрес", "official_status": "исходный статус"}
ALIASES = {
    "source": {"source", "источник", "канал"},
    "source_record_id": {"source_record_id", "external_id", "id", "id обращения", "id в источнике", "номер обращения", "идентификатор"},
    "date": {"date", "дата", "дата обращения", "created_at"},
    "text": {"text", "текст", "текст обращения", "обращение", "описание"},
    "address": {"address", "адрес", "адрес/район", "район", "территория"},
    "official_status": {"official_status", "status", "статус", "исходный статус"},
}


def fail(message: str, status: int = 400):
    raise HTTPException(status_code=status, detail=message)


def safe_filename(value: str | None) -> str:
    filename = (value or "import.csv").replace("\\", "/").rsplit("/", 1)[-1]
    return mask(filename).text[:180]


def read_file(content: bytes, filename: str) -> pd.DataFrame:
    if not content:
        fail("Файл пуст.")
    if len(content) > MAX_FILE_BYTES:
        fail("Максимальный размер файла — 5 МБ.", 413)
    extension = filename.lower().rsplit(".", 1)[-1]
    try:
        if extension == "csv":
            decoded = None
            for encoding in ("utf-8-sig", "cp1251"):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            if decoded is None or "\x00" in decoded:
                fail("Не удалось прочитать CSV. Сохраните файл в UTF-8.")
            try:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
                separator = dialect.delimiter
            except csv.Error:
                separator = ","
            reader = csv.reader(io.StringIO(decoded), delimiter=separator, strict=True)
            raw_rows = []
            row_numbers = []
            previous_line = 0
            for row in reader:
                current_line = previous_line + 1
                previous_line = reader.line_num
                if any(cell.strip() for cell in row):
                    raw_rows.append(row)
                    row_numbers.append(current_line)
                if len(raw_rows) > MAX_ROWS + 1:
                    fail(f"В одном файле допускается не более {MAX_ROWS} строк.")
            if not raw_rows:
                fail("В файле нет заголовков и данных.")
            headers = raw_rows[0]
            if len(raw_rows) - 1 > MAX_ROWS:
                fail(f"В одном файле допускается не более {MAX_ROWS} строк.")
            if any(len(row) != len(headers) for row in raw_rows[1:]):
                fail("CSV содержит строки с разным количеством столбцов. Проверьте разделитель и кавычки.")
            # Pandas is used for a consistent dataframe across CSV and XLSX.
            frame = pd.DataFrame(raw_rows[1:], columns=headers, index=row_numbers[1:], dtype=str)
        elif extension == "xlsx":
            try:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    members = archive.infolist()
                    if len(members) > 1000 or sum(item.file_size for item in members) > 40 * 1024 * 1024:
                        fail("Слишком большой распакованный XLSX.")
                    if any(item.file_size > 1024 * 1024 and item.file_size / max(item.compress_size, 1) > 150 for item in members):
                        fail("XLSX имеет недопустимую степень сжатия.")
                    if any("vbaProject" in item.filename for item in members):
                        fail("Файлы с макросами не поддерживаются.")
            except zipfile.BadZipFile:
                fail("Файл не является корректным XLSX.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    fail(f"Первый лист XLSX должен занимать не более {MAX_ROWS + 1} строк вместе с заголовком.")
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    fail(f"Допускается не более {MAX_COLUMNS} столбцов.")
                collected = []
                row_numbers = []
                for sheet_row_number, row in enumerate(sheet.iter_rows(), start=1):
                    if any(cell.data_type == "f" for cell in row):
                        fail("XLSX содержит формулы. Сохраните значения вместо формул перед импортом.")
                    values = [cell.value for cell in row]
                    if any(value is not None and str(value).strip() for value in values):
                        collected.append(values)
                        row_numbers.append(sheet_row_number)
                    if len(collected) > MAX_ROWS + 1:
                        fail(f"В одном файле допускается не более {MAX_ROWS} строк.")
                if not collected:
                    fail("В первом листе XLSX нет данных.")
                headers = [str(value).strip() if value is not None else "" for value in collected[0]]
                frame = pd.DataFrame(collected[1:], columns=headers, index=row_numbers[1:])
            finally:
                workbook.close()
        else:
            fail("Поддерживаются только файлы CSV и XLSX.", 415)
    except HTTPException:
        raise
    except Exception:
        # Parser messages may contain personal data from cells; never expose those messages.
        fail("Не удалось прочитать файл. Проверьте формат CSV/XLSX и отсутствие повреждений.")
    headers = [str(column).strip() for column in frame.columns]
    if not headers or len(headers) > MAX_COLUMNS:
        fail(f"Допускается от 1 до {MAX_COLUMNS} столбцов.")
    if any(not name or len(name) > 120 or UNSUPPORTED_CHARACTERS.search(name) for name in headers):
        fail("Заголовки столбцов должны быть непустыми и не длиннее 120 символов.")
    if len(set(headers)) != len(headers):
        fail("Заголовки столбцов повторяются. Сделайте их уникальными.")
    if any(mask(name).status != "clear" for name in headers):
        fail("В заголовках обнаружены вероятные персональные сведения. Используйте названия полей.")
    frame.columns = headers
    if frame.empty:
        fail("Файл содержит заголовки, но не содержит обращений.")
    return frame


def get_mapping(columns: list[str], requested: str | None) -> dict[str, str]:
    if requested:
        try:
            mapping = json.loads(requested)
        except (json.JSONDecodeError, TypeError):
            fail("Сопоставление столбцов должно быть корректным JSON.")
        if not isinstance(mapping, dict) or any(key not in FIELDS or not isinstance(value, (str, type(None))) for key, value in mapping.items()):
            fail("Некорректный формат сопоставления столбцов.")
        mapping = {key: value for key, value in mapping.items() if value}
        if any(value not in columns for value in mapping.values()):
            fail("Выбранный столбец отсутствует в файле.")
        if len(set(mapping.values())) != len(mapping):
            fail("Один столбец нельзя сопоставить нескольким полям.")
        return mapping
    return {field: next(column for column in columns if column.casefold() in aliases) for field, aliases in ALIASES.items() if any(column.casefold() in aliases for column in columns)}


def clean_value(value) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def normalize_date(value: str) -> str | None:
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(value, pattern).date().isoformat()
        except ValueError:
            pass
    if re.match(r"^\d{4}-\d{2}-\d{2}[T ]", value):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).date().isoformat()
        except ValueError:
            pass
    return None


def build_rows(frame: pd.DataFrame, mapping: dict[str, str], known: dict[tuple[str, str], str]) -> list[dict]:
    rows = []
    seen = dict(known)
    for number, raw in frame.iterrows():
        values = {field: clean_value(raw[mapping[field]]) if field in mapping else "" for field in FIELDS}
        errors = []
        warnings = []
        for field in REQUIRED:
            if not values[field]:
                errors.append(f"Обязательное поле «{LABELS[field]}» не заполнено.")
        for field, limit in (("source", 120), ("source_record_id", 160), ("text", MAX_TEXT), ("address", 500), ("official_status", 120)):
            if len(values[field]) > limit:
                errors.append(f"Поле «{LABELS[field]}» превышает {limit} символов.")
            if UNSUPPORTED_CHARACTERS.search(values[field]):
                errors.append(f"Поле «{LABELS[field]}» содержит недопустимые управляющие символы.")
        normalized_date = normalize_date(values["date"]) if values["date"] else None
        if values["date"] and not normalized_date:
            errors.append("Неверная дата. Используйте ГГГГ-ММ-ДД или ДД.ММ.ГГГГ.")
        values["date"] = normalized_date or ""
        if not values["address"]:
            values["address"] = "неизвестно"
            warnings.append("Адрес/территория не указаны: неизвестно.")
        if not values["official_status"]:
            values["official_status"] = "неизвестно"
        fingerprint = hashlib.sha256(json.dumps(values, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
        key = (values["source"], values["source_record_id"])
        language = detect_language(values["text"])
        statuses = []
        for field in ("text", "address", "source", "source_record_id", "official_status"):
            result = mask(values[field])
            values[field] = result.text
            statuses.append(result.status)
            warnings.extend(result.warnings)
            if field in ("source", "source_record_id", "official_status") and result.status != "clear":
                errors.append(f"Поле «{LABELS[field]}» содержит вероятные персональные сведения. Исправьте файл.")
        privacy_status = "review_required" if "review_required" in statuses else "masked" if "masked" in statuses else "clear"
        duplicate_status = "new"
        if not errors:
            if key in seen:
                duplicate_status = "duplicate" if seen[key] == fingerprint else "conflict"
                if duplicate_status == "conflict":
                    warnings.append("Этот источник и ID уже существуют с другим содержимым. Запись не будет перезаписана.")
            else:
                seen[key] = fingerprint
        # Only masked fields plus a digest are persisted; no original file or raw rows are retained.
        rows.append({"row_number": number, **values, "language": language, "privacy_status": privacy_status, "warnings": list(dict.fromkeys(warnings)), "errors": errors, "duplicate_status": duplicate_status, "original_hash": fingerprint})
    return rows
